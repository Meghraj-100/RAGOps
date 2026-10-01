"""Document ingestion service — extract text, chunk, embed, store."""

import os
import uuid
from io import BytesIO
from pathlib import Path

from sqlalchemy.ext.asyncio import AsyncSession

from app.core.config import get_settings
from app.models.models import Document, Chunk
from app.rag.embeddings import embed_texts
from app.observability.metrics import DOCUMENTS_INGESTED
import structlog

logger = structlog.get_logger()
settings = get_settings()

ALLOWED_TYPES = {
    "application/pdf": "pdf",
    "text/plain": "txt",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
}
ALLOWED_EXTENSIONS = {"pdf", "txt", "docx"}


def _extract_text_pdf(content: bytes) -> str:
    """Extract text from a PDF file."""
    from pypdf import PdfReader
    reader = PdfReader(BytesIO(content))
    pages = []
    for page in reader.pages:
        text = page.extract_text()
        if text:
            pages.append(text)
    return "\n\n".join(pages)


def _extract_text_docx(content: bytes) -> str:
    """Extract text from a DOCX file."""
    from docx import Document as DocxDocument
    doc = DocxDocument(BytesIO(content))
    return "\n\n".join(p.text for p in doc.paragraphs if p.text.strip())


def _extract_text_txt(content: bytes) -> str:
    """Extract text from a plain text file."""
    return content.decode("utf-8", errors="replace")


def extract_text(content: bytes, file_type: str) -> str:
    """Route to the appropriate text extractor."""
    extractors = {
        "pdf": _extract_text_pdf,
        "txt": _extract_text_txt,
        "docx": _extract_text_docx,
    }
    extractor = extractors.get(file_type)
    if not extractor:
        raise ValueError(f"Unsupported file type: {file_type}")
    return extractor(content)


def chunk_text(text: str, chunk_size: int | None = None, overlap: int | None = None) -> list[str]:
    """Split text into overlapping chunks of roughly chunk_size characters."""
    cs = chunk_size or settings.chunk_size
    ov = overlap or settings.chunk_overlap

    # Normalize whitespace
    text = " ".join(text.split())
    if not text:
        return []

    chunks = []
    start = 0
    while start < len(text):
        end = start + cs
        chunk = text[start:end]
        if chunk.strip():
            chunks.append(chunk.strip())
        start = end - ov

    return chunks


def get_file_extension(filename: str) -> str:
    """Get the lowercase file extension without the dot."""
    return Path(filename).suffix.lstrip(".").lower()


async def ingest_document(
    filename: str,
    content: bytes,
    content_type: str,
    db: AsyncSession,
) -> Document:
    """Full ingestion pipeline: validate → extract → chunk → embed → store."""

    file_type = get_file_extension(filename)
    if file_type not in ALLOWED_EXTENSIONS:
        raise ValueError(f"Unsupported file type: {file_type}. Allowed: {ALLOWED_EXTENSIONS}")

    # Create document record
    doc = Document(
        filename=filename,
        file_type=file_type,
        file_size=len(content),
        status="processing",
    )
    db.add(doc)
    await db.flush()

    try:
        # Extract text
        logger.info("extracting_text", filename=filename, file_type=file_type)
        text = extract_text(content, file_type)
        if not text.strip():
            raise ValueError("Document is empty — no text could be extracted")

        # Chunk
        logger.info("chunking_text", filename=filename)
        text_chunks = chunk_text(text)
        if not text_chunks:
            raise ValueError("No chunks produced from document")

        # Embed
        logger.info("embedding_chunks", filename=filename, count=len(text_chunks))
        embeddings = embed_texts(text_chunks)

        # Store chunks
        for i, (chunk_text_val, embedding) in enumerate(zip(text_chunks, embeddings)):
            chunk = Chunk(
                document_id=doc.id,
                content=chunk_text_val,
                chunk_index=i,
                embedding=embedding,
                metadata_={"source": filename, "chunk_index": i},
            )
            db.add(chunk)

        doc.status = "completed"
        doc.chunk_count = len(text_chunks)
        await db.flush()

        DOCUMENTS_INGESTED.inc()
        logger.info("ingestion_complete", filename=filename, chunks=len(text_chunks))
        return doc

    except Exception as e:
        doc.status = "failed"
        doc.error_message = str(e)[:500]
        await db.flush()
        logger.error("ingestion_failed", filename=filename, error=str(e))
        raise
